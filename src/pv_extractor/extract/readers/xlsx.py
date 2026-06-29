"""XlsxReader (D1): openpyxl read-only sheet -> cell grid, for client
valuation workbooks. One PageContent per worksheet (unit_label='sheet'); the
whole sheet becomes a single TableData grid (capped, see _MAX_ROWS/_COLS) and
its text joins the non-empty cells so targeting and the peek-verifier can
score sheets like pages."""

from __future__ import annotations

import io
from datetime import date, datetime
from pathlib import Path

from pv_extractor.models import DocFlag, DocumentContent, PageContent, TableData

from .base import DocumentReader

_MAX_ROWS = 500
_MAX_COLS = 64


def _cell_str(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == value.min.time() else value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    return text or None


class XlsxReader(DocumentReader):
    name = "xlsx"

    def summarize(self, path: str | Path, max_pages: int | None = None) -> DocumentContent:
        try:
            data = self._read_bytes(path)
        except OSError as exc:
            return self._empty(path, DocFlag.ACCESS_ERROR, f"{type(exc).__name__}: {exc}")
        if not data:
            return self._empty(path, DocFlag.CORRUPT_FILE, "zero-byte file")
        try:
            import openpyxl

            workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as exc:
            return self._empty(path, DocFlag.CORRUPT_FILE, f"{type(exc).__name__}: {exc}")

        pages: list[PageContent] = []
        try:
            for number, sheet in enumerate(workbook.worksheets, start=1):
                rows: list[list[str | None]] = []
                for row in sheet.iter_rows(max_row=_MAX_ROWS, max_col=_MAX_COLS, values_only=True):
                    cells = [_cell_str(value) for value in row]
                    while cells and cells[-1] is None:  # read-only mode pads to max_col
                        cells.pop()
                    rows.append(cells)
                while rows and not any(cell for cell in rows[-1]):
                    rows.pop()
                tables = (
                    [TableData(page_number=number, rows=rows, source=f"xlsx:{sheet.title}")]
                    if any(any(cell for cell in row) for row in rows)
                    else []
                )
                text = "\n".join(
                    " ".join(cell for cell in row if cell) for row in rows if any(cell for cell in row)
                )
                pages.append(
                    PageContent(
                        page_number=number,
                        text=text,
                        tables=tables,
                        text_char_count=len(text.strip()),
                        unit_label="sheet",
                        unit_name=sheet.title,
                    )
                )
        finally:
            workbook.close()
        shown = pages if max_pages is None else pages[:max_pages]
        return DocumentContent(
            file_path=str(path), reader=self.name, page_count=len(pages), pages=shown
        )

    def extract_tables(self, path: str | Path, page_numbers: list[int]) -> dict[int, list[TableData]]:
        content = self.summarize(path)
        wanted = set(page_numbers)
        return {
            page.page_number: page.tables
            for page in content.pages
            if page.page_number in wanted and page.tables
        }
